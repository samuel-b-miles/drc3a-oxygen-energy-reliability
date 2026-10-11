#!/usr/bin/env python3
"""Project only the two Ndosho tabs and keep manual and model-derived values distinct."""
import argparse,csv,gzip,hashlib,io,json,re
from pathlib import Path
from datetime import datetime
import openpyxl

def build(source:Path,out:Path):
    out.mkdir(parents=True,exist_ok=True)
    raw=openpyxl.load_workbook(source,data_only=False)
    cache=openpyxl.load_workbook(source,data_only=True)
    records=[];inventory=[]
    for title in ['ChCbNdosho 2025','ChCbNdosho 2023']:
        s=raw[title]; c=cache[title];dates=[]
        for row in s:
            for cell in row:
                if cell.value is None:continue
                value=cell.value
                if isinstance(value,str):
                    value=re.sub(r'https?://\S+','[controlled source link omitted]',value)
                    value=re.sub(r'Ch Cb Ndosho','N001',value,flags=re.I)
                cached=c[cell.coordinate].value
                if isinstance(cached,str):
                    cached=re.sub(r'https?://\S+','[controlled source link omitted]',cached)
                    cached=re.sub(r'Ch Cb Ndosho','N001',cached,flags=re.I)
                if cell.column==1 and isinstance(cached,datetime):dates.append(cached.isoformat())
                records.append({'collection_id':'N001','source_tab':title,'cell':cell.coordinate,
                                'source_value':str(value),'cached_value':str(cached) if cached is not None else '',
                                'source_type':cell.data_type,
                                'evidence_layer':'model-derived prediction' if 14<=cell.column<=17 else 'manual or spreadsheet-derived register'})
        inventory.append({'source_tab':title,'rows':s.max_row,'columns':s.max_column,
                          'date_column_min_as_stored':min(dates) if dates else None,
                          'date_column_max_as_stored':max(dates) if dates else None,
                          'date_status':'2023 dates and copied header need day/month reconciliation' if '2023' in title else 'stored 2024–2025 dates; continuations require day carry-forward'})
    with (out/'source_cell_projection.csv.gz').open('wb') as stream,gzip.GzipFile(filename='',mode='wb',fileobj=stream,mtime=0) as gz:
        with io.TextIOWrapper(gz,encoding='utf-8',newline='') as f:
            w=csv.DictWriter(f,fieldnames=list(records[0]));w.writeheader();w.writerows(records)
    s=cache['ChCbNdosho 2025'];event=[]
    for date,rows in [('2025-01-26',[50]),('2025-01-27',[51]),('2025-01-28',[52]),('2025-01-29',[53,54])]:
        event.append({'collection_id':'N001','date_local':date,'source_tab':'ChCbNdosho 2025',
                      'source_cells_runtime':';'.join(f'K{r}' for r in rows),
                      'logged_generator_minutes':sum(s.cell(r,11).value for r in rows),
                      'source_cells_fuel':';'.join(f'F{r}' for r in rows),
                      'fuel_column_liters_recorded':sum((s.cell(r,6).value or 0) for r in rows),
                      'fuel_interpretation':'entered under Fuel Consumption; purchase/dispensing/burn distinction requires original-register audit',
                      'generator_groups':';'.join(str(int(s.cell(r,2).value)) for r in rows)})
    with (out/'event_register_observations.csv').open('w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(event[0]));w.writeheader();w.writerows(event)
    manual=sum(x['logged_generator_minutes'] for x in event)/60
    model=sum(s.cell(r,17).value for r in range(37,41))
    # Carry dates through continuation rows, without filling absent calendar days.
    current_date=None; january_minutes=0; january_dates=set()
    for row in range(8,90):
        value=s.cell(row,1).value
        if isinstance(value,datetime):current_date=value.date()
        minutes=s.cell(row,11).value
        if current_date and str(current_date).startswith('2025-01') and isinstance(minutes,(int,float)):
            january_minutes+=minutes; january_dates.add(str(current_date))
    report={'source_url':'https://docs.google.com/spreadsheets/d/1Ru5n5YFWkVHveDQ8NX89NhDA17ry_PydR3vUo-MhAgY/edit',
            'source_modified_utc':'2025-07-28T04:40:56.270Z','retrieved_date_utc':'2026-10-11',
            'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'source_credit':'Daniel Steven Doxazo Sitompul: confirmed 2025 logbook spreadsheet entry/edit contributor from July 9, 2025 cell-level Sheets history; sole transcription and model authorship not established.',
            'digitization_provenance':'digitization_provenance.json',
            'drc3a_inclusion_status':'included provisionally by study lead, 2026-10-11; source-photo and event-boundary audits pending',
            'january_runtime_sensitivity':{
                'source_cells':'ChCbNdosho 2025!A8:A89,K8:K89; date carry-forward for continuation rows',
                'logged_hours_in_January':january_minutes/60,
                'distinct_January_dates_with_numeric_runtime':len(january_dates),
                'calendar_dates_with_no_numeric_runtime_entry':sorted(set(f'2025-01-{d:02d}' for d in range(1,32))-january_dates),
                'four_day_logged_hours':manual,
                'four_day_share_of_recorded_January_hours_pct':100*manual/(january_minutes/60),
                'logged_January_hours_excluding_four_days':january_minutes/60-manual,
                'model_vs_logged_four_day_difference_pct':100*(model-manual)/manual,
                'interpretation':'Recorded generator runtime only; missing dates are not zero. This is not total facility energy, metered fuel burn, or a routine-consumption estimate.'},
            'tabs':inventory,'event_manual_hours':manual,'event_model_derived_hours':round(model,8),
            'event_fuel_liters_entered':sum(x['fuel_column_liters_recorded'] for x in event),
            'model_cells':'ChCbNdosho 2025!N37:Q40','manual_cells':'ChCbNdosho 2025!A50:L54',
            'limitations':['Model-derived generator hours are not an independent runtime measurement. Model code and training data are not linked in this workbook.',
                           'Manual and model-derived event hours disagree and must remain separate.',
                           '2023 tab has a copied 2024–2025 timeline header and conflicting day/month interpretations.',
                           'Dashes mark continuation, not zero runtime. Fuel entries are not independent metered fuel-burn observations.']}
    (out/'source_manifest.json').write_text(json.dumps(report,indent=2)+'\n')
    raw.close();cache.close();print(json.dumps(report,indent=2))

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--source',type=Path,default=Path('data/raw/private/ndosho/generator_2025/source.xlsx'))
    p.add_argument('--out',type=Path,default=Path('data/collections/ndosho/generator_register'))
    a=p.parse_args();build(a.source,a.out)
