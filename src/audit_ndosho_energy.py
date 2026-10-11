#!/usr/bin/env python3
"""Read-only audit of measured total load and historical ledger proxies."""
from pathlib import Path
import hashlib
import json
import numpy as np
import pandas as pd
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'data/collections/ndosho/energy_audit'

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def build():
    OUT.mkdir(parents=True, exist_ok=True)
    source = ROOT / 'data/raw/private/ndosho/hopmeter/source_repo/data/total_load.csv'
    workbook = ROOT / 'data/raw/private/ndosho/historical/source.xlsx'
    x = pd.read_csv(source)
    units = x['Total Power'].str.extract(r'^\s*([+-]?[\d.]+)\s*(kW|W)\s*$')
    x['power_kW'] = units[0].astype(float) * np.where(units[1].eq('kW'), 1, .001)
    x['time'] = pd.to_datetime(x.Time, errors='raise')
    if x.time.duplicated().any() or not np.isfinite(x.power_kW).all() or x.power_kW.lt(0).any():
        raise ValueError('Resolve duplicates or invalid power before integration')
    x = x.sort_values('time')
    delta = x.time.diff().dt.total_seconds()
    # Do not bridge a missing minute or extrapolate over larger gaps.
    x['covered_seconds'] = delta.where(delta.gt(0) & delta.le(60), 0)
    x['energy_kWh'] = ((x.power_kW + x.power_kW.shift()) / 2 * x.covered_seconds / 3600).fillna(0)
    daily = x.groupby(x.time.dt.date).agg(energy_kWh=('energy_kWh', 'sum'),
                                         covered_seconds=('covered_seconds', 'sum'), rows=('time', 'size'))
    daily['coverage_fraction'] = daily.covered_seconds / 86400
    daily['include_95pct'] = daily.coverage_fraction.ge(.95)
    daily.to_csv(OUT / 'hopmeter_daily_energy.csv', index_label='date_source_clock')
    good = daily[daily.include_95pct]
    hourly = x.set_index('time').power_kW.resample('h').mean()
    profile = hourly.groupby(hourly.index.hour).agg(['mean', 'median', 'count'])
    profile.to_csv(OUT / 'hopmeter_hourly_profile.csv', index_label='hour_source_clock')
    book = openpyxl.load_workbook(workbook, data_only=True)
    ledger = {}
    for title, label in [('comparisons - usage', 'energy_proxy'), ('comparisons - purchases', 'cost_proxy')]:
        sheet = book[title]
        rows = []
        for r in range(2, sheet.max_row + 1):
            date, week, grid, generator = [sheet.cell(r, c).value for c in range(1, 5)]
            if isinstance(grid, (int, float)) and isinstance(generator, (int, float)):
                rows.append(dict(source_row=r, date_cached=date, week=week, grid=grid, generator=generator))
        frame = pd.DataFrame(rows)
        frame['date_cached'] = pd.to_datetime(frame.date_cached)
        frame.to_csv(OUT / f'{label}_matched_weeks.csv', index=False)
        annual = frame.groupby(frame.date_cached.dt.year).agg(weeks=('grid', 'size'),
                       grid_total=('grid', 'sum'), generator_total=('generator', 'sum'),
                       grid_mean_week=('grid', 'mean'), generator_mean_week=('generator', 'mean'))
        annual['generator_share'] = annual.generator_total / (annual.grid_total + annual.generator_total)
        annual.to_csv(OUT / f'{label}_annual.csv', index_label='year_cached')
        grid, generator = float(frame.grid.sum()), float(frame.generator.sum())
        ledger[label] = dict(source_tab=title, source_range=f'A{frame.source_row.min()}:D{frame.source_row.max()}',
            matched_weeks=len(frame), first_week_label=str(frame.date_cached.min()), last_week_label=str(frame.date_cached.max()),
            grid_total=grid, generator_total=generator, generator_share=generator/(grid+generator),
            grid_share=grid/(grid+generator), mean_total_per_week=(grid+generator)/len(frame))
    report = dict(
        audit_date='2026-10-11', source_files_unchanged=True,
        total_load=dict(source_url='https://github.com/astroash11/drc3/blob/c21987d8c0cb3d3653d6aad485ee32661e3a5565/data/total_load.csv',
            source_sha256=digest(source), rows=len(x), first_time=str(x.time.min()), last_time=str(x.time.max()),
            timezone='unverified source clock', total_load_placement='source whole-facility label; physical deployment not independently verified',
            method='Sort timestamps; parse explicit W/kW; trapezoidal power integration on intervals <=60 seconds; no gap filling or prorating; assign intervals to endpoint source-clock date.',
            days_with_95pct_coverage=len(good), mean_daily_energy_kWh=float(good.energy_kWh.mean()),
            median_daily_energy_kWh=float(good.energy_kWh.median()),
            daily_p10_kWh=float(good.energy_kWh.quantile(.1)), daily_p90_kWh=float(good.energy_kWh.quantile(.9)),
            mean_observed_load_kW=float(x.power_kW.mean()),
            mean_daily_energy_at_99pct_coverage_kWh=float(daily.loc[daily.coverage_fraction.ge(.99), 'energy_kWh'].mean()),
            cached_workbook_load_profile='comparisons - usage!AA76:AB99; total generator usage!U2',
            cached_profile_implied_daily_kWh=float(book['total generator usage']['U2'].value)/7),
        ledger=ledger, workbook_sha256=digest(workbook),
        energy_assumptions=dict(source_cells='total generator usage!P2:Q3', generator_I_kW=64, generator_II_kW=16,
            formula='80*0.8 and 20*0.8; fixed assumed output multiplied by logged/inferred runtime',
            status='Not measured generator electrical output. Generator-II calculation warnings, date alignment and missing-to-zero formulas remain unresolved.'),
        cost_assumptions=dict(source_cells='comparisons - purchases!C14:D124; generators!R:T,V1',
            grid_tariff_USD_per_kWh=.2, generator_purchase_units='Source comment: units are probably dollars',
            status='Conditional bookkeeping proxy, not verified USD expenditure. Grid purchases differ from energy used; fuel inventory, unit price, maintenance and battery/solar supply are not reconciled.'),
        existing_notebook_issue=dict(source_url='https://github.com/astroash11/drc3/blob/c21987d8c0cb3d3653d6aad485ee32661e3a5565/full_facility_box_plot.ipynb',
            issue='Sums 60 kW readings without dividing by 60 and labels as kWh. Also groups unsorted source rows into positional blocks. Those energy results are not used.'),
        period_limits='2023 total-load baseline and 2020-2022 ledger shares are different periods. Do not apply historical energy shares to 2023 measured kWh/day or describe them as current post-PSA demand. January 2025 logged generator hours are not whole-facility kWh.',
        script_sha256=digest(Path(__file__))
    )
    book.close()
    (OUT / 'audit_summary.json').write_text(json.dumps(report, indent=2) + '\n')
    print(json.dumps({'mean_daily_kWh':report['total_load']['mean_daily_energy_kWh'], 'ledger':ledger}, indent=2))

if __name__ == '__main__':
    build()
